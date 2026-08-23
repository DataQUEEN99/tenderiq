import io
import fitz
from typing import BinaryIO
from app.shared.exceptions.custom_exceptions import AIProcessingException


class TextExtractor:
    def __init__(self):
        self.supported_types = {
            "application/pdf": self._extract_pdf,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": self._extract_docx,
            "text/plain": self._extract_text,
            "text/markdown": self._extract_text,
        }

    async def extract(self, file_path: str, mime_type: str) -> str:
        extractor = self.supported_types.get(mime_type)
        if not extractor:
            raise AIProcessingException(f"Unsupported file type: {mime_type}")
        return await extractor(file_path)

    async def extract_from_bytes(self, content: bytes, mime_type: str) -> str:
        extractor = self.supported_types.get(mime_type)
        if not extractor:
            raise AIProcessingException(f"Unsupported file type: {mime_type}")

        if mime_type == "application/pdf":
            return await self._extract_pdf_from_bytes(content)
        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return await self._extract_docx_from_bytes(content)
        else:
            return content.decode("utf-8")

    async def _extract_pdf(self, file_path: str) -> str:
        try:
            doc = fitz.open(file_path)
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            doc.close()
            return "\n\n".join(text_parts)
        except Exception as e:
            raise AIProcessingException(f"PDF extraction failed: {str(e)}")

    async def _extract_pdf_from_bytes(self, content: bytes) -> str:
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            text_parts = []
            for page in doc:
                text_parts.append(page.get_text())
            doc.close()
            return "\n\n".join(text_parts)
        except Exception as e:
            raise AIProcessingException(f"PDF extraction failed: {str(e)}")

    async def _extract_docx(self, file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            return "\n\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            raise AIProcessingException(f"DOCX extraction failed: {str(e)}")

    async def _extract_docx_from_bytes(self, content: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            raise AIProcessingException(f"DOCX extraction failed: {str(e)}")

    async def _extract_text(self, file_path: str) -> str:
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except Exception as e:
            raise AIProcessingException(f"Text extraction failed: {str(e)}")