from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.dialects.postgresql import UUID
from typing import List, Tuple, Optional
from app.modules.export.repository import ExportRepository
from app.modules.export.models import Export, ExportFormat, ExportStatus
from app.modules.export.schemas import ExportCreate, ExportResponse
from app.modules.proposal.models import Proposal
from app.shared.exceptions.custom_exceptions import ResourceNotFoundException
from app.shared.logging.logger import app_logger
from app.shared.storage.storage_service import StorageService
from datetime import datetime
from uuid import uuid4


class ExportService:
    def __init__(self, session: AsyncSession):
        self.repo = ExportRepository(session)
        self.session = session
        self.storage = StorageService()

    async def create_export(self, organization_id: UUID, created_by: UUID, request: ExportCreate) -> ExportResponse:
        from uuid import UUID as UUIDType
        proposal_id = UUIDType(request.proposal_id)

        from app.modules.proposal.service import ProposalService
        proposal_service = ProposalService(self.session)
        proposal_detail = await proposal_service.get_proposal_with_sections(proposal_id, organization_id)

        export = Export(
            proposal_id=proposal_id,
            export_format=request.export_format,
            organization_id=organization_id,
            created_by=created_by,
        )
        export = await self.repo.create(export)

        export.status = ExportStatus.PROCESSING
        await self.repo.update(export)

        try:
            file_path = await self._generate_export(proposal_detail, request.export_format, export.export_id, organization_id)
            export.file_path = file_path
            export.status = ExportStatus.COMPLETED
        except Exception as e:
            export.status = ExportStatus.FAILED
            export.error_message = str(e)
            app_logger.error(f"Export failed", extra={"export_id": str(export.export_id), "error": str(e)})

        export = await self.repo.update(export)
        return ExportResponse.model_validate(export)

    async def _generate_export(self, proposal, export_format: ExportFormat, export_id: UUID, organization_id: UUID) -> str:
        sections = getattr(proposal, "sections", [])
        content = self._build_proposal_content(proposal, sections)

        if export_format == ExportFormat.PDF:
            file_path = await self._export_pdf(proposal, content, export_id, organization_id)
        elif export_format == ExportFormat.DOCX:
            file_path = await self._export_docx(proposal, content, export_id, organization_id)
        else:
            file_path = await self._export_html(proposal, content, export_id, organization_id)

        return file_path

    def _build_proposal_content(self, proposal, sections) -> str:
        parts = [f"# {proposal.title}", ""]
        if getattr(proposal, "executive_summary", None):
            parts.append(proposal.executive_summary)
            parts.append("")
        for section in sections:
            is_deleted = getattr(section, "is_deleted", False)
            if not is_deleted:
                parts.append(f"## {section.title}")
                if getattr(section, "content", None):
                    parts.append(section.content)
                parts.append("")
        return "\n".join(parts)

    async def _export_pdf(self, proposal, content: str, export_id: UUID, organization_id: UUID = None) -> str:
        from fpdf import FPDF
        import os

        org_id = organization_id or getattr(proposal, "organization_id", None)
        proposal_id = getattr(proposal, "id", None)

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)

        for line in content.split("\n"):
            if line.startswith("# "):
                pdf.set_font("Helvetica", "B", 18)
                pdf.cell(0, 10, line[2:], ln=True)
                pdf.set_font("Helvetica", size=12)
            elif line.startswith("## "):
                pdf.set_font("Helvetica", "B", 14)
                pdf.cell(0, 8, line[3:], ln=True)
                pdf.set_font("Helvetica", size=12)
            elif line.strip():
                pdf.multi_cell(0, 6, line)
            else:
                pdf.ln(4)

        filename = f"proposal_{proposal_id}_{export_id}.pdf"
        path = f"organizations/{org_id}/exports/{filename}"

        pdf_output = pdf.output(dest="S").encode("latin-1")
        await self.storage.save(path, pdf_output)

        return path

    async def _export_docx(self, proposal, content: str, export_id: UUID, organization_id: UUID = None) -> str:
        from docx import Document
        import io

        org_id = organization_id or getattr(proposal, "organization_id", None)
        proposal_id = getattr(proposal, "id", None)
        proposal_title = getattr(proposal, "title", "Proposal")

        doc = Document()
        doc.add_heading(proposal_title, 0)

        for line in content.split("\n"):
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.strip():
                doc.add_paragraph(line)

        filename = f"proposal_{proposal_id}_{export_id}.docx"
        path = f"organizations/{org_id}/exports/{filename}"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        await self.storage.save(path, buffer.read())

        return path

    async def _export_html(self, proposal, content: str, export_id: UUID, organization_id: UUID = None) -> str:
        org_id = organization_id or getattr(proposal, "organization_id", None)
        proposal_id = getattr(proposal, "id", None)
        proposal_title = getattr(proposal, "title", "Proposal")

        html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{proposal_title}</title>
    <style>
        body {{ font-family: Arial, sans-serif; max-width: 800px; margin: 0 auto; padding: 40px; }}
        h1 {{ color: #1a365d; border-bottom: 2px solid #2b6cb0; padding-bottom: 10px; }}
        h2 {{ color: #2b6cb0; margin-top: 30px; }}
        p {{ line-height: 1.6; }}
    </style>
</head>
<body>
"""

        for line in content.split("\n"):
            if line.startswith("# "):
                html += f"<h1>{line[2:]}</h1>\n"
            elif line.startswith("## "):
                html += f"<h2>{line[3:]}</h2>\n"
            elif line.strip():
                html += f"<p>{line}</p>\n"
            else:
                html += "<br>\n"

        html += "</body></html>"

        filename = f"proposal_{proposal_id}_{export_id}.html"
        path = f"organizations/{org_id}/exports/{filename}"
        await self.storage.save(path, html.encode("utf-8"))

        return path

    async def get_export(self, export_id: UUID, organization_id: UUID) -> ExportResponse:
        export = await self.repo.get_by_id(export_id, organization_id)
        if not export or export.is_deleted:
            raise ResourceNotFoundException("Export not found")
        return ExportResponse.model_validate(export)

    async def list_exports(self, organization_id: UUID, page: int = 1, page_size: int = 20) -> Tuple[List[ExportResponse], int]:
        items, total = await self.repo.list(organization_id, page, page_size)
        return [ExportResponse.model_validate(item) for item in items], total

    async def delete_export(self, export_id: UUID, organization_id: UUID) -> None:
        export = await self.repo.get_by_id(export_id, organization_id)
        if not export:
            raise ResourceNotFoundException("Export not found")

        if export.file_path:
            try:
                await self.storage.delete(export.file_path)
            except Exception as e:
                app_logger.warning(f"Failed to delete export file", extra={"error": str(e)})

        await self.repo.soft_delete(export_id, organization_id)
        app_logger.info(f"Export deleted", extra={"export_id": str(export_id)})

    async def download_export(self, export_id: UUID, organization_id: UUID) -> bytes:
        export = await self.repo.get_by_id(export_id, organization_id)
        if not export or export.is_deleted:
            raise ResourceNotFoundException("Export not found")

        if export.status != ExportStatus.COMPLETED or not export.file_path:
            raise ResourceNotFoundException("Export file not available")

        return await self.storage.read(export.file_path)